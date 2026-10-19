"""Employees (SCR-10 … SCR-15): registry, masking, import, duplicates, deactivation."""

from __future__ import annotations

import csv
import io
from datetime import date, timedelta

from pydantic import BaseModel, Field, field_validator

from .. import db as dbm
from ..errors import bad_request, conflict, not_found, unprocessable
from ..http import Ctx, page_params, paged, route
from ..security import crypto, passwords
from ..services import gallery, notify, tenancy
from ..util import dumps, loads, new_id, now_iso

NAME_RE = r"^[A-Za-zА-Яа-яЁёЎўҚқҒғҲҳʻʼ'`\- ]+$"


class EmployeeIn(BaseModel):
    personnel_no: str = Field(min_length=1, max_length=20, pattern=r"^[A-Za-z0-9\-]+$")
    last_name: str = Field(min_length=1, max_length=60, pattern=NAME_RE)
    first_name: str = Field(min_length=1, max_length=60, pattern=NAME_RE)
    middle_name: str | None = Field(default=None, max_length=60, pattern=NAME_RE)
    pinfl: str | None = Field(default=None, pattern=r"^\d{14}$")
    birth_date: date | None = None
    phone: str | None = Field(default=None, pattern=r"^\+\d{9,15}$")
    site_id: str
    department_id: str
    position_id: str | None = None
    team: str | None = Field(default=None, max_length=10)
    hire_date: date
    access_point_ids: list[str] | None = None
    card_no: str | None = Field(default=None, max_length=32)
    pin: str | None = None
    language: str = Field(default="uz", pattern="^(uz|ru|en)$")

    @field_validator("personnel_no")
    @classmethod
    def upper(cls, v: str) -> str:
        return v.upper()

    @field_validator("birth_date")
    @classmethod
    def age(cls, v):
        if v is None:
            return v
        today = date.today()
        years = (today - v).days / 365.25
        if v > today or not 14 <= years <= 80:
            raise ValueError("age must be between 14 and 80")
        return v

    @field_validator("hire_date")
    @classmethod
    def hire(cls, v):
        if v > date.today() + timedelta(days=60):
            raise ValueError("hire date can be at most 60 days ahead")
        return v


class EmployeePatch(BaseModel):
    version: int
    last_name: str | None = Field(default=None, max_length=60, pattern=NAME_RE)
    first_name: str | None = Field(default=None, max_length=60, pattern=NAME_RE)
    middle_name: str | None = Field(default=None, max_length=60, pattern=NAME_RE)
    pinfl: str | None = Field(default=None, pattern=r"^\d{14}$")
    birth_date: date | None = None
    phone: str | None = Field(default=None, pattern=r"^\+\d{9,15}$")
    department_id: str | None = None
    position_id: str | None = None
    team: str | None = Field(default=None, max_length=10)
    access_point_ids: list[str] | None = None
    card_no: str | None = Field(default=None, max_length=32)
    pin: str | None = None
    language: str | None = Field(default=None, pattern="^(uz|ru|en)$")


class DeactivateIn(BaseModel):
    effective_date: date
    reason: str = Field(pattern="^(dismissed|transferred|other)$")
    comment: str | None = Field(default=None, max_length=500)


class BulkIn(BaseModel):
    ids: list[str] = Field(min_length=1, max_length=500)
    action: str = Field(pattern="^(move_department|deactivate|invite_enrollment)$")
    department_id: str | None = None
    reason: str | None = None


class MergeIn(BaseModel):
    primary_id: str
    comment: str = Field(min_length=10, max_length=500)


SENSITIVE = ("pinfl_last4", "phone", "birth_date")


def view(row: dict, sensitive: bool) -> dict:
    out = {k: v for k, v in row.items() if k not in ("pinfl_enc", "pinfl_hash", "pin_hash")}
    out["has_pin"] = bool(row.get("pin_hash"))
    if not sensitive:
        for k in SENSITIVE:
            out.pop(k, None)
        out.pop("card_no", None)
    else:
        out["pinfl_masked"] = f"••••••••••{row['pinfl_last4']}" if row.get("pinfl_last4") else None
        out.pop("pinfl_last4", None)
    return out


def get_employee(ctx: Ctx, emp_id: str, perm: str = "employee.read") -> dict:
    row = dbm.one(ctx.conn, "SELECT * FROM employees WHERE id = ? AND tenant_id = ?", (emp_id, ctx.tenant_id))
    if not row:
        raise not_found("Employee")
    ctx.require_employee(perm, row)
    return row


def _check_refs(ctx: Ctx, site_id: str, department_id: str, position_id: str | None) -> None:
    site = dbm.one(ctx.conn, "SELECT id FROM sites WHERE id = ? AND tenant_id = ?", (site_id, ctx.tenant_id))
    if not site:
        raise unprocessable("Unknown site", "SITE_UNKNOWN", field="site_id")
    ctx.require_site("employee.write", site_id)
    dept = dbm.one(ctx.conn, "SELECT site_id, status FROM departments WHERE id = ? AND tenant_id = ?",
                   (department_id, ctx.tenant_id))
    if not dept or dept["site_id"] != site_id or dept["status"] != "active":
        raise unprocessable("Department does not belong to the site", "DEPARTMENT_INVALID", field="department_id")
    if position_id and not dbm.one(ctx.conn, "SELECT 1 AS x FROM positions WHERE id = ? AND tenant_id = ?",
                                   (position_id, ctx.tenant_id)):
        raise unprocessable("Unknown position", "POSITION_UNKNOWN", field="position_id")


def _secrets(ctx: Ctx, data: dict, pinfl: str | None, pin: str | None, card_no: str | None, emp_id: str | None):
    if pinfl:
        data["pinfl_enc"] = crypto.seal(ctx.dek(), pinfl.encode(), b"pinfl")
        data["pinfl_hash"] = crypto.keyed_hash(ctx.dek(), pinfl)
        data["pinfl_last4"] = pinfl[-4:]
    if pin:
        problems = passwords.pin_problems(pin)
        if problems:
            raise unprocessable("PIN is too simple or malformed", "PIN_WEAK", field="pin", problems=problems)
        data["pin_hash"] = passwords.hash_secret(pin)
    if card_no is not None:
        card_no = card_no.strip() or None
        if card_no:
            other = dbm.one(ctx.conn, "SELECT id, personnel_no FROM employees WHERE tenant_id = ? AND card_no = ?",
                            (ctx.tenant_id, card_no))
            if other and other["id"] != emp_id:
                raise conflict(f"Card is already assigned to {other['personnel_no']}", "CARD_TAKEN", field="card_no")
        data["card_no"] = card_no


def find_duplicates(ctx: Ctx, emp: dict) -> list[dict]:
    found = []
    if emp.get("pinfl_hash"):
        for o in dbm.all_(ctx.conn, "SELECT id, personnel_no, last_name, first_name FROM employees WHERE tenant_id = ? "
                                    "AND pinfl_hash = ? AND id <> ?", (ctx.tenant_id, emp["pinfl_hash"], emp["id"])):
            found.append({**o, "reason": "pinfl"})
    if emp.get("birth_date"):
        for o in dbm.all_(ctx.conn, "SELECT id, personnel_no, last_name, first_name FROM employees WHERE tenant_id = ? "
                                    "AND lower(last_name) = lower(?) AND lower(first_name) = lower(?) AND birth_date = ? "
                                    "AND id <> ?", (ctx.tenant_id, emp["last_name"], emp["first_name"],
                                                    emp["birth_date"], emp["id"])):
            if not any(f["id"] == o["id"] for f in found):
                found.append({**o, "reason": "name_birth"})
    for f in found:
        a, b = sorted([emp["id"], f["id"]])
        if not dbm.one(ctx.conn, "SELECT 1 AS x FROM duplicate_candidates WHERE tenant_id = ? AND employee_a = ? AND "
                                 "employee_b = ? AND reason = ?", (ctx.tenant_id, a, b, f["reason"])):
            dbm.insert(ctx.conn, "duplicate_candidates", {"id": new_id(), "tenant_id": ctx.tenant_id, "employee_a": a,
                                                          "employee_b": b, "reason": f["reason"], "score": None,
                                                          "status": "open", "decided_by": None, "decided_at": None,
                                                          "created_at": now_iso()})
    return found


def _set_access_points(ctx: Ctx, emp_id: str, site_id: str, ap_ids: list[str] | None) -> None:
    if ap_ids is None:
        ap_ids = [r["id"] for r in dbm.all_(ctx.conn, "SELECT id FROM access_points WHERE tenant_id = ? AND site_id = ? "
                                                      "AND status = 'active'", (ctx.tenant_id, site_id))]
    else:
        valid = {r["id"] for r in dbm.all_(ctx.conn, "SELECT id FROM access_points WHERE tenant_id = ? AND site_id = ?",
                                           (ctx.tenant_id, site_id))}
        bad = [a for a in ap_ids if a not in valid]
        if bad:
            raise unprocessable("Unknown access points for this site", "ACCESS_POINT_INVALID", ids=bad)
        if not ap_ids:
            raise unprocessable("At least one access point is required", "ACCESS_POINT_REQUIRED")
    ctx.conn.execute("DELETE FROM employee_access_points WHERE employee_id = ?", (emp_id,))
    for ap in ap_ids:
        ctx.conn.execute("INSERT INTO employee_access_points VALUES (?,?,?)", (ctx.tenant_id, emp_id, ap))


def create_employee_row(ctx: Ctx, body: EmployeeIn, source: str = "manual") -> tuple[dict, list[dict]]:
    _check_refs(ctx, body.site_id, body.department_id, body.position_id)
    tenancy.check_limit(ctx.conn, ctx.tenant_id, "employees")
    if dbm.one(ctx.conn, "SELECT 1 AS x FROM employees WHERE tenant_id = ? AND site_id = ? AND personnel_no = ?",
               (ctx.tenant_id, body.site_id, body.personnel_no)):
        other = dbm.one(ctx.conn, "SELECT last_name, first_name FROM employees WHERE tenant_id = ? AND site_id = ? AND "
                                  "personnel_no = ?", (ctx.tenant_id, body.site_id, body.personnel_no))
        raise conflict(f"Personnel number {body.personnel_no} is already used ({other['last_name']} "
                       f"{other['first_name'][:1]}.)", "PERSONNEL_NO_TAKEN", field="personnel_no")
    now = now_iso()
    emp = {"id": new_id(), "tenant_id": ctx.tenant_id, "site_id": body.site_id, "personnel_no": body.personnel_no,
           "last_name": body.last_name.strip(), "first_name": body.first_name.strip(),
           "middle_name": (body.middle_name or "").strip() or None,
           "birth_date": body.birth_date.isoformat() if body.birth_date else None, "phone": body.phone,
           "department_id": body.department_id, "position_id": body.position_id, "team": body.team,
           "hire_date": body.hire_date.isoformat(), "status": "active", "language": body.language,
           "biometric_status": "none", "version": 1, "created_at": now, "updated_at": now}
    _secrets(ctx, emp, body.pinfl, body.pin, body.card_no, None)
    dbm.insert(ctx.conn, "employees", emp)
    dbm.insert(ctx.conn, "employment_assignments", {"id": new_id(), "tenant_id": ctx.tenant_id, "employee_id": emp["id"],
                                                    "department_id": body.department_id, "position_id": body.position_id,
                                                    "team": body.team, "valid_from": body.hire_date.isoformat(),
                                                    "valid_to": None, "created_by": ctx.actor, "created_at": now})
    _set_access_points(ctx, emp["id"], body.site_id, body.access_point_ids)
    dups = find_duplicates(ctx, emp)
    ctx.audit("employee.created", "employee", emp["id"], after={**emp, "source": source})
    return emp, dups


@route("GET", "/api/v1/employees", perm="employee.read")
def list_employees(ctx: Ctx):
    scope_sql, params = ctx.employee_filter("employee.read")
    where = ["e.tenant_id = ?", scope_sql]
    params = [ctx.tenant_id, *params]
    tab = ctx.q("tab", "active")
    if tab == "active":
        where.append("e.status = 'active'")
    elif tab == "deactivated":
        where.append("e.status = 'deactivated'")
    elif tab == "no_biometric":
        where.append("e.status = 'active' AND e.biometric_status IN ('none','revoked','declined')")
    q = (ctx.q("q", "") or "").strip()
    if q:
        if len(q) < 2:
            raise unprocessable("Search needs at least 2 characters", "QUERY_TOO_SHORT")
        like = f"%{q.lower()}%"
        where.append("(lower(e.last_name) LIKE ? OR lower(e.first_name) LIKE ? OR lower(e.personnel_no) LIKE ? "
                     "OR (e.phone IS NOT NULL AND substr(e.phone, -4) = ?))")
        params += [like, like, like, q[-4:] if q.isdigit() else "~"]
    for col, name in (("site_id", "site_id"), ("department_id", "department_id"), ("position_id", "position_id"),
                      ("team", "team"), ("biometric_status", "biometric_status")):
        vals = ctx.qlist(name)
        if vals:
            where.append(f"e.{col} IN ({','.join('?' * len(vals))})")
            params += vals
    if ctx.q("hired_from"):
        where.append("e.hire_date >= ?")
        params.append(ctx.q("hired_from"))
    if ctx.q("hired_to"):
        where.append("e.hire_date <= ?")
        params.append(ctx.q("hired_to"))
    sort = {"name": "e.last_name, e.first_name", "personnel_no": "e.personnel_no", "hire_date": "e.hire_date DESC",
            "department": "d.name, e.last_name"}.get(ctx.q("sort", "name"), "e.last_name, e.first_name")
    limit, offset = page_params(ctx)
    sql_where = " AND ".join(where)
    total = dbm.scalar(ctx.conn, f"SELECT COUNT(*) FROM employees e LEFT JOIN departments d ON d.id = e.department_id "
                                 f"WHERE {sql_where}", params)
    rows = dbm.all_(ctx.conn, f"SELECT e.*, d.name AS department_name, p.name AS position_name, s.name AS site_name, "
                              f"(SELECT max(captured_at) FROM attendance_events ev WHERE ev.employee_id = e.id "
                              f"AND ev.result IN ('accepted','review')) AS last_event_at "
                              f"FROM employees e LEFT JOIN departments d ON d.id = e.department_id "
                              f"LEFT JOIN positions p ON p.id = e.position_id LEFT JOIN sites s ON s.id = e.site_id "
                              f"WHERE {sql_where} ORDER BY {sort} LIMIT ? OFFSET ?", [*params, limit, offset])
    sensitive = ctx.can("employee.read_sensitive")
    counts = {}
    base_where = ["e.tenant_id = ?", scope_sql]
    base_params = [ctx.tenant_id, *ctx.employee_filter("employee.read")[1]]
    for name, cond in (("active", "e.status = 'active'"), ("deactivated", "e.status = 'deactivated'"),
                       ("no_biometric", "e.status = 'active' AND e.biometric_status IN ('none','revoked','declined')"),
                       ("all", "1=1")):
        counts[name] = dbm.scalar(ctx.conn, f"SELECT COUNT(*) FROM employees e WHERE {' AND '.join(base_where)} AND {cond}",
                                  base_params)
    return {**paged([view(r, sensitive) for r in rows], total, limit, offset), "counts": counts,
            "sensitive": sensitive}


@route("GET", "/api/v1/employees/{id}", perm="employee.read")
def get_one(ctx: Ctx):
    emp = get_employee(ctx, ctx.path["id"])
    sensitive = ctx.can("employee.read_sensitive")
    assignments = dbm.all_(ctx.conn, "SELECT a.*, d.name AS department_name, p.name AS position_name "
                                     "FROM employment_assignments a LEFT JOIN departments d ON d.id = a.department_id "
                                     "LEFT JOIN positions p ON p.id = a.position_id WHERE a.employee_id = ? "
                                     "ORDER BY a.valid_from DESC", (emp["id"],))
    aps = dbm.all_(ctx.conn, "SELECT ap.id, ap.name, ap.direction FROM employee_access_points eap JOIN access_points ap "
                             "ON ap.id = eap.access_point_id WHERE eap.employee_id = ?", (emp["id"],))
    profile = dbm.one(ctx.conn, "SELECT id, status, model_version, quality_score, liveness_score, frames, gallery_version, "
                                "created_at FROM biometric_profiles WHERE employee_id = ? AND status <> 'revoked' "
                                "ORDER BY created_at DESC LIMIT 1", (emp["id"],))
    consent = dbm.one(ctx.conn, "SELECT decision, policy_version, method, created_at FROM biometric_consents "
                                "WHERE employee_id = ? ORDER BY created_at DESC LIMIT 1", (emp["id"],))
    last_ok = dbm.scalar(ctx.conn, "SELECT max(captured_at) FROM attendance_events WHERE employee_id = ? AND method = 'face' "
                                   "AND result = 'accepted'", (emp["id"],))
    dept = dbm.one(ctx.conn, "SELECT name FROM departments WHERE id = ?", (emp["department_id"],))
    pos = dbm.one(ctx.conn, "SELECT name FROM positions WHERE id = ?", (emp["position_id"],)) if emp["position_id"] else None
    site = dbm.one(ctx.conn, "SELECT name, timezone FROM sites WHERE id = ?", (emp["site_id"],))
    return {**view(emp, sensitive), "department_name": dept["name"] if dept else None,
            "position_name": pos["name"] if pos else None, "site_name": site["name"], "timezone": site["timezone"],
            "assignments": assignments, "access_points": aps,
            "biometric": {"profile": profile, "consent": consent, "last_face_match_at": last_ok}}


@route("POST", "/api/v1/employees", perm="employee.write", billing="grow")
def create(ctx: Ctx, body: EmployeeIn):
    emp, dups = create_employee_row(ctx, body)
    out = view(emp, ctx.can("employee.read_sensitive"))
    return 201, {**out, "duplicate_warnings": dups}


@route("PATCH", "/api/v1/employees/{id}", perm="employee.write", billing="write")
def patch(ctx: Ctx, body: EmployeePatch):
    emp = get_employee(ctx, ctx.path["id"], "employee.write")
    if body.version != emp["version"]:
        raise conflict("Data was changed by another user. Reload and try again.", "VERSION_CONFLICT",
                       current_version=emp["version"])
    data = body.model_dump(exclude_unset=True, exclude={"version", "pinfl", "pin", "card_no", "access_point_ids"})
    if "birth_date" in data and data["birth_date"]:
        data["birth_date"] = data["birth_date"].isoformat()
    _secrets(ctx, data, body.pinfl, body.pin, body.card_no if "card_no" in body.model_fields_set else None, emp["id"])
    dept_changed = "department_id" in data and data["department_id"] != emp["department_id"]
    if dept_changed or "position_id" in data or "team" in data:
        _check_refs(ctx, emp["site_id"], data.get("department_id", emp["department_id"]),
                    data.get("position_id", emp["position_id"]))
        today = date.today().isoformat()
        ctx.conn.execute("UPDATE employment_assignments SET valid_to = ? WHERE employee_id = ? AND valid_to IS NULL",
                         (today, emp["id"]))
        dbm.insert(ctx.conn, "employment_assignments", {
            "id": new_id(), "tenant_id": ctx.tenant_id, "employee_id": emp["id"],
            "department_id": data.get("department_id", emp["department_id"]),
            "position_id": data.get("position_id", emp["position_id"]), "team": data.get("team", emp["team"]),
            "valid_from": today, "valid_to": None, "created_by": ctx.actor, "created_at": now_iso()})
    data["version"] = emp["version"] + 1
    data["updated_at"] = now_iso()
    dbm.update(ctx.conn, "employees", {"id": emp["id"]}, data)
    if body.access_point_ids is not None:
        _set_access_points(ctx, emp["id"], emp["site_id"], body.access_point_ids)
        gallery.bump(ctx.conn, ctx.tenant_id)
    if "card_no" in data or "pin_hash" in data:
        gallery.bump(ctx.conn, ctx.tenant_id)
    ctx.audit("employee.updated", "employee", emp["id"], before=emp, after=data)
    updated = dbm.one(ctx.conn, "SELECT * FROM employees WHERE id = ?", (emp["id"],))
    return view(updated, ctx.can("employee.read_sensitive"))


def deactivate_employee(ctx: Ctx, emp: dict, effective: date, reason: str, comment: str | None) -> None:
    open_corr = dbm.all_(ctx.conn, "SELECT id, work_date FROM corrections WHERE employee_id = ? AND status = 'pending'",
                         (emp["id"],))
    if open_corr:
        raise conflict("Employee has pending corrections; resolve them first", "PENDING_CORRECTIONS",
                       corrections=open_corr)
    if effective > date.today() + timedelta(days=90):
        raise unprocessable("Effective date must be within 90 days", "DATE_TOO_FAR")
    data = {"status": "deactivated", "deactivated_at": effective.isoformat(),
            "deactivation_reason": f"{reason}: {comment or ''}".strip(": "), "version": emp["version"] + 1,
            "updated_at": now_iso()}
    dbm.update(ctx.conn, "employees", {"id": emp["id"]}, data)
    ctx.conn.execute("UPDATE users SET status = 'blocked' WHERE employee_id = ?", (emp["id"],))
    ctx.conn.execute("UPDATE sessions SET revoked_at = ? WHERE user_id IN (SELECT id FROM users WHERE employee_id = ?) "
                     "AND revoked_at IS NULL", (now_iso(), emp["id"]))
    gallery.bump(ctx.conn, ctx.tenant_id)
    ctx.audit("employee.deactivated", "employee", emp["id"], before={"status": emp["status"]}, after=data)


@route("POST", "/api/v1/employees/{id}/deactivate", perm="employee.deactivate")
def deactivate(ctx: Ctx, body: DeactivateIn):
    emp = get_employee(ctx, ctx.path["id"], "employee.deactivate")
    if emp["status"] == "deactivated":
        raise conflict("Employee is already deactivated", "ALREADY_DEACTIVATED")
    deactivate_employee(ctx, emp, body.effective_date, body.reason, body.comment)
    return {"ok": True, "status": "deactivated"}


@route("POST", "/api/v1/employees/{id}/reactivate", perm="employee.deactivate", billing="grow")
def reactivate(ctx: Ctx):
    emp = get_employee(ctx, ctx.path["id"], "employee.deactivate")
    if emp["status"] != "deactivated":
        raise conflict("Employee is active", "NOT_DEACTIVATED")
    tenancy.check_limit(ctx.conn, ctx.tenant_id, "employees")
    dbm.update(ctx.conn, "employees", {"id": emp["id"]}, {"status": "active", "deactivated_at": None,
                                                         "deactivation_reason": None, "version": emp["version"] + 1,
                                                         "updated_at": now_iso()})
    ctx.conn.execute("UPDATE users SET status = 'active' WHERE employee_id = ? AND password_hash IS NOT NULL", (emp["id"],))
    gallery.bump(ctx.conn, ctx.tenant_id)
    ctx.audit("employee.reactivated", "employee", emp["id"])
    needs_enrollment = emp["biometric_status"] in ("none", "revoked")
    return {"ok": True, "needs_enrollment": needs_enrollment}


@route("POST", "/api/v1/employees/bulk", perm="employee.write")
def bulk(ctx: Ctx, body: BulkIn):
    done, skipped = 0, []
    for emp_id in body.ids:
        emp = dbm.one(ctx.conn, "SELECT * FROM employees WHERE id = ? AND tenant_id = ?", (emp_id, ctx.tenant_id))
        if not emp or not ctx.can_employee("employee.write", emp):
            skipped.append({"id": emp_id, "reason": "not_found_or_out_of_scope"})
            continue
        if body.action == "move_department":
            if not body.department_id:
                raise unprocessable("department_id is required", "DEPARTMENT_REQUIRED")
            _check_refs(ctx, emp["site_id"], body.department_id, emp["position_id"])
            today = date.today().isoformat()
            ctx.conn.execute("UPDATE employment_assignments SET valid_to = ? WHERE employee_id = ? AND valid_to IS NULL",
                             (today, emp["id"]))
            dbm.insert(ctx.conn, "employment_assignments", {
                "id": new_id(), "tenant_id": ctx.tenant_id, "employee_id": emp["id"],
                "department_id": body.department_id, "position_id": emp["position_id"], "team": emp["team"],
                "valid_from": today, "valid_to": None, "created_by": ctx.actor, "created_at": now_iso()})
            dbm.update(ctx.conn, "employees", {"id": emp["id"]}, {"department_id": body.department_id,
                                                                 "version": emp["version"] + 1, "updated_at": now_iso()})
            ctx.audit("employee.moved", "employee", emp["id"], before={"department_id": emp["department_id"]},
                      after={"department_id": body.department_id})
        elif body.action == "deactivate":
            ctx.require("employee.deactivate")
            if emp["status"] == "deactivated":
                skipped.append({"id": emp_id, "reason": "already_deactivated"})
                continue
            try:
                deactivate_employee(ctx, emp, date.today(), body.reason or "other", "bulk")
            except Exception as e:  # noqa: BLE001
                skipped.append({"id": emp_id, "reason": getattr(e, "code", "error")})
                continue
        elif body.action == "invite_enrollment":
            notify.notify_employee(ctx.conn, ctx.tenant_id, emp["id"], "enrollment_invite",
                                   "Yuzni ro'yxatdan o'tkazish",
                                   "HR bo'limiga yuzni ro'yxatdan o'tkazish uchun keling (rozilik ixtiyoriy).")
        done += 1
    return {"done": done, "skipped": skipped}


# ----------------------------------------------------------------- import (SCR-13)

IMPORT_COLUMNS = ["personnel_no", "last_name", "first_name", "middle_name", "pinfl", "birth_date", "phone",
                  "site_code", "department_code", "position", "team", "hire_date", "card_no", "language"]


def _parse_file(filename: str, content: bytes) -> list[dict]:
    name = filename.lower()
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig")
        dialect = csv.Sniffer().sniff(text.splitlines()[0]) if text else csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        return [{(k or "").strip().lower(): (v or "").strip() for k, v in row.items()} for row in reader]
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(h or "").strip().lower() for h in rows[0]]
        out = []
        for r in rows[1:]:
            if all(v in (None, "") for v in r):
                continue
            item = {}
            for k, v in zip(header, r):
                if hasattr(v, "date") and callable(getattr(v, "date")):
                    v = v.date().isoformat()
                item[k] = "" if v is None else str(v).strip()
            out.append(item)
        return out
    raise unprocessable("Only .csv and .xlsx files are supported", "FILE_TYPE")


@route("GET", "/api/v1/employees/import-template", perm="employee.import")
def import_template(ctx: Ctx):
    from starlette.responses import Response
    import xlsxwriter
    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True})
    ws = wb.add_worksheet("employees")
    bold = wb.add_format({"bold": True, "bg_color": "#E8EEF8"})
    for i, col in enumerate(IMPORT_COLUMNS):
        ws.write(0, i, col, bold)
        ws.set_column(i, i, 16)
    ws.write_row(1, 0, ["A-1001", "Karimov", "Akmal", "Olimovich", "", "1990-05-12", "+998901234567", "Z1",
                        "SEX-1", "Operator", "A", "2026-01-10", "", "uz"])
    wb.close()
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename=fras-employees-template.xlsx"})


@route("POST", "/api/v1/employees/imports", perm="employee.import", form=True, billing="grow")
def import_preview(ctx: Ctx, form: dict):
    files = form["files"]
    if "file" not in files:
        raise bad_request("Attach a file in field 'file'", "FILE_REQUIRED")
    filename, content = files["file"]
    if len(content) > 10 * 1024 * 1024:
        raise unprocessable("File is larger than 10 MB", "FILE_TOO_LARGE")
    rows = _parse_file(filename, content)
    if len(rows) > 10_000:
        raise unprocessable("At most 10 000 rows per import", "TOO_MANY_ROWS")
    sites = {r["code"].upper(): r for r in dbm.all_(ctx.conn, "SELECT * FROM sites WHERE tenant_id = ?", (ctx.tenant_id,))}
    depts = {(r["site_id"], r["code"].upper()): r for r in dbm.all_(ctx.conn, "SELECT * FROM departments WHERE tenant_id = ?",
                                                                   (ctx.tenant_id,))}
    positions = {r["name"].lower(): r for r in dbm.all_(ctx.conn, "SELECT * FROM positions WHERE tenant_id = ?",
                                                        (ctx.tenant_id,))}
    preview, seen = [], set()
    summary = {"new": 0, "update": 0, "error": 0, "duplicate": 0}
    for i, raw in enumerate(rows, start=2):
        errors = []
        site = sites.get((raw.get("site_code") or "").upper())
        if not site:
            errors.append("site_code: unknown")
        dept = depts.get((site["id"], (raw.get("department_code") or "").upper())) if site else None
        if site and not dept:
            errors.append("department_code: unknown")
        pos = positions.get((raw.get("position") or "").lower()) if raw.get("position") else None
        payload = {
            "personnel_no": raw.get("personnel_no"), "last_name": raw.get("last_name"),
            "first_name": raw.get("first_name"), "middle_name": raw.get("middle_name") or None,
            "pinfl": raw.get("pinfl") or None, "birth_date": raw.get("birth_date") or None,
            "phone": raw.get("phone") or None, "site_id": site["id"] if site else "-",
            "department_id": dept["id"] if dept else "-", "position_id": pos["id"] if pos else None,
            "team": raw.get("team") or None, "hire_date": raw.get("hire_date") or date.today().isoformat(),
            "card_no": raw.get("card_no") or None, "language": raw.get("language") or "uz",
        }
        try:
            EmployeeIn.model_validate(payload)
        except Exception as e:  # pydantic ValidationError
            for err in getattr(e, "errors", lambda: [])():
                errors.append(f"{'.'.join(str(x) for x in err['loc'])}: {err['msg']}")
        key = (payload["site_id"], (payload["personnel_no"] or "").upper())
        status = "error" if errors else "new"
        if not errors:
            if key in seen:
                status, errors = "error", ["duplicate row in file"]
            elif dbm.one(ctx.conn, "SELECT 1 AS x FROM employees WHERE tenant_id = ? AND site_id = ? AND personnel_no = ?",
                         (ctx.tenant_id, key[0], key[1])):
                status = "update"
        seen.add(key)
        summary["error" if status == "error" else status] += 1
        preview.append({"row": i, "status": status, "errors": errors, "data": payload})
    batch = {"id": new_id(), "tenant_id": ctx.tenant_id, "kind": "employees", "filename": filename,
             "file_hash": crypto.sha256_hex(content), "status": "preview", "summary": dumps(summary),
             "rows": dumps(preview), "created_by": ctx.actor, "created_at": now_iso(), "applied_at": None}
    dbm.insert(ctx.conn, "import_batches", batch)
    return 201, {"id": batch["id"], "summary": summary, "rows": preview[:200], "total_rows": len(preview)}


@route("POST", "/api/v1/employees/imports/{id}/apply", perm="employee.import", billing="grow")
def import_apply(ctx: Ctx):
    batch = dbm.one(ctx.conn, "SELECT * FROM import_batches WHERE id = ? AND tenant_id = ?", (ctx.path["id"], ctx.tenant_id))
    if not batch:
        raise not_found("Import")
    if batch["status"] != "preview":
        raise conflict("Import was already applied", "IMPORT_ALREADY_APPLIED")
    created = updated = 0
    errors = []
    for item in loads(batch["rows"], []):
        if item["status"] == "error":
            continue
        data = item["data"]
        existing = dbm.one(ctx.conn, "SELECT * FROM employees WHERE tenant_id = ? AND site_id = ? AND personnel_no = ?",
                           (ctx.tenant_id, data["site_id"], data["personnel_no"].upper()))
        try:
            if existing:
                changes = {k: data[k] for k in ("last_name", "first_name", "middle_name", "phone", "team", "language",
                                                "department_id", "position_id", "birth_date") if data.get(k)}
                _secrets(ctx, changes, data.get("pinfl"), None, data.get("card_no"), existing["id"])
                if any(existing.get(k) != v for k, v in changes.items()):
                    changes["version"] = existing["version"] + 1
                    changes["updated_at"] = now_iso()
                    dbm.update(ctx.conn, "employees", {"id": existing["id"]}, changes)
                    ctx.audit("employee.updated", "employee", existing["id"], before=existing,
                              after={**changes, "source": "import"})
                    updated += 1
            else:
                create_employee_row(ctx, EmployeeIn.model_validate(data), source="import")
                created += 1
        except Exception as e:  # noqa: BLE001 - report per-row failures
            errors.append({"row": item["row"], "error": getattr(e, "message", str(e))})
    dbm.update(ctx.conn, "import_batches", {"id": batch["id"]},
               {"status": "applied", "applied_at": now_iso(),
                "summary": dumps({**loads(batch["summary"]), "created": created, "updated": updated,
                                  "apply_errors": len(errors)})})
    ctx.audit("employee.import_applied", "import", batch["id"], after={"created": created, "updated": updated,
                                                                       "errors": len(errors)})
    return {"created": created, "updated": updated, "errors": errors}


@route("GET", "/api/v1/employees/imports/{id}", perm="employee.import")
def import_get(ctx: Ctx):
    batch = dbm.one(ctx.conn, "SELECT * FROM import_batches WHERE id = ? AND tenant_id = ?", (ctx.path["id"], ctx.tenant_id))
    if not batch:
        raise not_found("Import")
    return {**batch, "summary": loads(batch["summary"]), "rows": loads(batch["rows"])}


# ----------------------------------------------------------------- duplicates (SCR-14)

@route("GET", "/api/v1/employees/duplicates", perm="employee.merge")
def duplicates(ctx: Ctx):
    rows = dbm.all_(ctx.conn, "SELECT dc.*, a.personnel_no AS a_no, a.last_name AS a_last, a.first_name AS a_first, "
                              "a.status AS a_status, b.personnel_no AS b_no, b.last_name AS b_last, "
                              "b.first_name AS b_first, b.status AS b_status FROM duplicate_candidates dc "
                              "JOIN employees a ON a.id = dc.employee_a JOIN employees b ON b.id = dc.employee_b "
                              "WHERE dc.tenant_id = ? AND dc.status = 'open' ORDER BY dc.created_at DESC",
                    (ctx.tenant_id,))
    if not (ctx.can("biometric.review") or ctx.can("biometric.read")):
        rows = [r for r in rows if r["reason"] != "face"]
    return {"items": rows}


@route("POST", "/api/v1/employees/duplicates/{id}/merge", perm="employee.merge")
def merge(ctx: Ctx, body: MergeIn):
    dc = dbm.one(ctx.conn, "SELECT * FROM duplicate_candidates WHERE id = ? AND tenant_id = ?", (ctx.path["id"], ctx.tenant_id))
    if not dc or dc["status"] != "open":
        raise not_found("Duplicate candidate")
    if body.primary_id not in (dc["employee_a"], dc["employee_b"]):
        raise unprocessable("primary_id must be one of the pair", "PRIMARY_INVALID")
    secondary_id = dc["employee_b"] if body.primary_id == dc["employee_a"] else dc["employee_a"]
    primary = dbm.one(ctx.conn, "SELECT * FROM employees WHERE id = ?", (body.primary_id,))
    secondary = dbm.one(ctx.conn, "SELECT * FROM employees WHERE id = ?", (secondary_id,))
    fill = {k: secondary[k] for k in ("middle_name", "phone", "birth_date", "pinfl_enc", "pinfl_hash", "pinfl_last4")
            if not primary.get(k) and secondary.get(k)}
    if fill:
        fill["version"] = primary["version"] + 1
        dbm.update(ctx.conn, "employees", {"id": primary["id"]}, fill)
    # Raw events stay attached to the original record (append-only); the secondary record is closed.
    dbm.update(ctx.conn, "employees", {"id": secondary["id"]},
               {"status": "deactivated", "deactivated_at": date.today().isoformat(),
                "deactivation_reason": f"merged into {primary['personnel_no']}: {body.comment}",
                "version": secondary["version"] + 1, "updated_at": now_iso()})
    dbm.update(ctx.conn, "duplicate_candidates", {"id": dc["id"]},
               {"status": "merged", "decided_by": ctx.actor, "decided_at": now_iso()})
    gallery.bump(ctx.conn, ctx.tenant_id)
    ctx.audit("employee.merged", "employee", primary["id"], before={"secondary": secondary_id},
              after={"comment": body.comment})
    return {"ok": True, "primary_id": primary["id"], "deactivated_id": secondary_id}


@route("POST", "/api/v1/employees/duplicates/{id}/distinct", perm="employee.merge")
def distinct(ctx: Ctx):
    dc = dbm.one(ctx.conn, "SELECT * FROM duplicate_candidates WHERE id = ? AND tenant_id = ?", (ctx.path["id"], ctx.tenant_id))
    if not dc:
        raise not_found("Duplicate candidate")
    dbm.update(ctx.conn, "duplicate_candidates", {"id": dc["id"]},
               {"status": "distinct", "decided_by": ctx.actor, "decided_at": now_iso()})
    ctx.audit("employee.duplicate_dismissed", "duplicate", dc["id"])
    return {"ok": True}


@route("GET", "/api/v1/employees/{id}/history", perm="employee.read")
def history(ctx: Ctx):
    emp = get_employee(ctx, ctx.path["id"])
    if not ctx.can("audit.read") and not ctx.can("employee.write"):
        return {"items": []}
    rows = dbm.all_(ctx.conn, "SELECT a.seq, a.action, a.actor_type, a.actor_id, u.name AS actor_name, a.created_at, "
                              "a.before_json, a.after_json FROM audit_logs a LEFT JOIN users u ON u.id = a.actor_id "
                              "WHERE a.tenant_id = ? AND a.object_type = 'employee' AND a.object_id = ? "
                              "ORDER BY a.seq DESC LIMIT 200", (ctx.tenant_id, emp["id"]))
    return {"items": rows}
